"""Importar proveedores y clientes desde Excel (.xlsx) o CSV.

Misma regla que el alta manual: solo el nombre es obligatorio. Telefono, RFC,
domicilio, correo y condiciones se cargan si vienen y se omiten (con aviso) si
traen un valor invalido. Los que ya existen no se tocan ni se duplican.

Flujo en dos pasos con el mismo archivo: primero "Revisar" (no guarda nada) y
despues "Importar" (guarda solo los nuevos).
"""
import csv
import io
import re
import unicodedata
from datetime import date, datetime
from decimal import Decimal

from flask import jsonify, redirect, render_template, request, send_file, session, url_for

MAX_FILAS = 2000
MAX_BYTES = 3 * 1024 * 1024
MAX_FILAS_VISTA = 500
CORREO_REGEX = re.compile(r"^[^@\s]+@[^@\s]+\.[^@\s]+$")
RFC_REGEX = re.compile(r"^[A-ZÑ&]{3,4}\d{6}[A-Z0-9]{3}$")


class ErrorImportar(ValueError):
    """Problema con el archivo completo (no con una fila)."""


def _clave(texto):
    t = unicodedata.normalize("NFKD", str(texto or "").casefold())
    t = "".join(c for c in t if not unicodedata.combining(c))
    return re.sub(r"[^a-z0-9]+", " ", t).strip()


# Nombre de columna que aceptamos (sin acentos ni mayusculas) -> campo interno.
ENCABEZADOS = {
    "comunes": {
        "nombre": ["nombre", "razon social", "nombre o razon social", "nombre completo", "proveedor",
                   "cliente", "empresa", "nombre del proveedor", "nombre del cliente"],
        "telefono": ["telefono", "tel", "celular", "whatsapp", "movil", "telefono celular", "numero de telefono"],
        "correo": ["correo", "email", "e mail", "mail", "correo electronico"],
        "rfc": ["rfc"],
        "direccion": ["direccion", "domicilio", "calle", "calle y numero", "calle numero"],
    },
    "proveedores": {
        "nombre_comercial": ["nombre comercial", "marca"],
        "contacto": ["contacto", "nombre del contacto", "persona de contacto", "atencion"],
        "colonia": ["colonia"],
        "ciudad": ["ciudad", "municipio"],
        "codigo_postal": ["cp", "c p", "codigo postal"],
        "condicion_pago": ["condicion de pago", "condiciones de pago", "condiciones", "forma de pago", "pago"],
        "tiempo_entrega": ["dias de entrega", "tiempo de entrega", "entrega", "dias entrega", "lead time"],
        "tipo": ["tipo", "tipo de proveedor", "giro"],
        "notas": ["notas", "comentarios", "observaciones"],
    },
    "clientes": {
        "apellidos": ["apellidos", "apellido"],
        "colonia": ["colonia"],
        "referencias": ["referencias", "referencia", "entre calles"],
        "zona": ["zona", "region"],
    },
}

PLANTILLAS = {
    "proveedores": {
        "columnas": ["Nombre", "Teléfono", "Correo", "RFC", "Contacto", "Dirección", "Colonia", "Ciudad",
                     "C.P.", "Condición de pago", "Días de entrega", "Tipo", "Notas"],
        "ejemplo": ["Química del Sureste", "998 123 4567", "ventas@quimicasureste.mx", "", "Ana Pérez",
                    "Av. Tulum 123", "Centro", "Cancún", "77500", "Contado", "3", "Materia prima", ""],
        "obligatorio": "Nombre",
    },
    "clientes": {
        "columnas": ["Nombre", "Apellidos", "Teléfono", "Correo", "RFC", "Dirección", "Colonia", "Referencias", "Zona"],
        "ejemplo": ["María", "López", "998 765 4321", "", "", "Calle 12 mz 4 lt 8", "SM 200", "Portón negro", "Cancun"],
        "obligatorio": "Nombre",
    },
}


def _alias(tipo):
    mapa = {}
    for grupo in (ENCABEZADOS["comunes"], ENCABEZADOS[tipo]):
        for campo, nombres in grupo.items():
            for n in nombres:
                mapa.setdefault(_clave(n), campo)
    return mapa


# ---------------------------------------------------------------- lectura

def _texto_celda(v):
    if v is None:
        return ""
    if isinstance(v, bool):
        return "Sí" if v else "No"
    if isinstance(v, float) and v.is_integer():
        v = int(v)
    if isinstance(v, (datetime, date)):
        return v.isoformat()
    return str(v).strip()


def _filas_xlsx(contenido):
    try:
        from openpyxl import load_workbook
        libro = load_workbook(io.BytesIO(contenido), read_only=True, data_only=True)
    except Exception as e:  # zip roto, no es Excel, etc.
        raise ErrorImportar("No pude abrir el archivo como Excel (.xlsx). Guárdalo de nuevo como .xlsx o .csv.") from e
    try:
        hoja = libro.worksheets[0]
        filas = []
        for fila in hoja.iter_rows(values_only=True):
            filas.append([_texto_celda(c) for c in fila])
            if len(filas) > MAX_FILAS + 30:
                break
        return filas
    finally:
        libro.close()


def _filas_csv(contenido):
    for codificacion in ("utf-8-sig", "cp1252"):
        try:
            texto = contenido.decode(codificacion)
            break
        except UnicodeDecodeError:
            continue
    else:
        raise ErrorImportar("No pude leer el texto del archivo. Guárdalo como CSV UTF-8 o como .xlsx.")
    primera = next((l for l in texto.splitlines() if l.strip()), "")
    delimitador = max([",", ";", "\t"], key=primera.count)
    lector = csv.reader(io.StringIO(texto), delimiter=delimitador)
    filas = []
    for fila in lector:
        filas.append([c.strip() for c in fila])
        if len(filas) > MAX_FILAS + 30:
            break
    return filas


def leer_archivo(nombre, contenido, tipo):
    """Devuelve (filas, ignoradas): filas = [(numero_de_fila, {campo: texto})]."""
    if len(contenido) > MAX_BYTES:
        raise ErrorImportar("El archivo pesa demasiado (máximo 3 MB). Divídelo en partes.")
    ext = (nombre or "").lower().rsplit(".", 1)[-1] if "." in (nombre or "") else ""
    if ext in ("xlsx", "xlsm"):
        crudas = _filas_xlsx(contenido)
    elif ext in ("csv", "txt"):
        crudas = _filas_csv(contenido)
    elif ext == "xls":
        raise ErrorImportar("Los .xls antiguos no se pueden leer. En Excel usa Guardar como → Libro de Excel (.xlsx) o CSV.")
    else:
        raise ErrorImportar("Sube un archivo de Excel (.xlsx) o CSV.")

    alias = _alias(tipo)
    for i, fila in enumerate(crudas[:10]):
        columnas = [(j, alias.get(_clave(c))) for j, c in enumerate(fila)]
        if any(campo == "nombre" for _, campo in columnas):
            encabezado = i
            break
    else:
        raise ErrorImportar(
            "No encontré la columna «Nombre» en las primeras filas. Descarga la plantilla o renombra la columna "
            "(también acepto «Razón social», «Proveedor» o «Cliente»)."
        )
    ignoradas = [c for j, c in enumerate(crudas[encabezado]) if c and columnas[j][1] is None]
    filas = []
    for k, fila in enumerate(crudas[encabezado + 1:], start=encabezado + 2):
        datos = {}
        for j, campo in columnas:
            if campo and j < len(fila) and fila[j] and not datos.get(campo):
                datos[campo] = fila[j]
        if datos:
            filas.append((k, datos))
    if len(filas) > MAX_FILAS:
        raise ErrorImportar(f"El archivo trae más de {MAX_FILAS} filas. Divídelo en partes.")
    return filas, ignoradas


# ------------------------------------------------------------- utilidades

def _recortar(valor, maximo, etiqueta, avisos):
    valor = (valor or "").strip()
    if len(valor) > maximo:
        avisos.append(f"{etiqueta} muy largo: se cortó a {maximo} caracteres")
        return valor[:maximo]
    return valor


def _telefono_clave(tel):
    digitos = re.sub(r"\D", "", tel or "")
    return digitos[-10:] if len(digitos) >= 7 else (tel or "").strip().casefold()


def _numero(texto):
    m = re.search(r"\d+(?:[.,]\d+)?", texto or "")
    return Decimal(m.group(0).replace(",", ".")) if m else None


def _fila(numero, nombre, estado, motivo="", avisos=None, datos=None):
    return {"fila": numero, "nombre": nombre, "estado": estado, "motivo": motivo,
            "avisos": avisos or [], "datos": datos or {}}


def _resumen(filas, ignoradas):
    return {
        "nuevos": sum(1 for f in filas if f["estado"] == "nuevo"),
        "existentes": sum(1 for f in filas if f["estado"] == "existe"),
        "omitidas": sum(1 for f in filas if f["estado"] in ("omitida", "repetida")),
        "con_avisos": sum(1 for f in filas if f["estado"] == "nuevo" and f["avisos"]),
        "total": len(filas),
        "columnas_ignoradas": ignoradas,
    }


def _plantilla_xlsx(tipo):
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill
    from openpyxl.utils import get_column_letter

    p = PLANTILLAS[tipo]
    libro = Workbook()
    datos = libro.active
    datos.title = "Datos"
    datos.append(p["columnas"])
    for celda in datos[1]:
        celda.font = Font(bold=True, color="FFFFFF")
        celda.fill = PatternFill("solid", fgColor="1F3A5F")
    for j, titulo in enumerate(p["columnas"], start=1):
        datos.column_dimensions[get_column_letter(j)].width = max(14, len(titulo) + 4)
    datos.freeze_panes = "A2"
    ejemplo = libro.create_sheet("Ejemplo")
    ejemplo.append(p["columnas"])
    ejemplo.append(p["ejemplo"])
    for celda in ejemplo[1]:
        celda.font = Font(bold=True)
    for j, titulo in enumerate(p["columnas"], start=1):
        ejemplo.column_dimensions[get_column_letter(j)].width = max(14, len(titulo) + 4)
    ayuda = libro.create_sheet("Instrucciones")
    for linea in (
        "Llena la hoja «Datos» (la primera). Una fila por " + ("proveedor." if tipo == "proveedores" else "cliente."),
        "Lo único obligatorio es el Nombre. Todo lo demás es opcional: déjalo vacío si no lo tienes.",
        "Puedes borrar columnas que no uses o reordenarlas; solo conserva el encabezado «Nombre».",
        "Los que ya existen en el sistema no se duplican ni se modifican.",
        "La hoja «Ejemplo» solo muestra cómo se ve una fila llena; no se importa.",
    ):
        ayuda.append([linea])
    ayuda.column_dimensions["A"].width = 110
    salida = io.BytesIO()
    libro.save(salida)
    salida.seek(0)
    return salida


# ----------------------------------------------------------------- registro

def registrar_importacion(app, db, Proveedor, Contacto, Bitacora, ClienteCRM, ClienteDomicilio,
                          admin_required_json, usuario_actual, siguiente_codigo, rfc_marcador, rfc_visible):

    def analizar_proveedores(filas):
        nombres = set()
        rfcs = set()
        for p in Proveedor.query.all():
            nombres.add(_clave(p.razon_social))
            if p.nombre_comercial:
                nombres.add(_clave(p.nombre_comercial))
            if rfc_visible(p.rfc):
                rfcs.add(p.rfc)
        vistos_nombre, vistos_rfc = {}, {}
        resultado = []
        for numero, d in filas:
            avisos = []
            nombre = _recortar(d.get("nombre"), 180, "Nombre", avisos)
            if not nombre:
                resultado.append(_fila(numero, "", "omitida", "Sin nombre"))
                continue
            rfc = re.sub(r"\s+", "", d.get("rfc", "")).upper()
            if rfc and not RFC_REGEX.match(rfc):
                avisos.append(f"RFC «{rfc}» no válido: se importa sin RFC")
                rfc = ""
            clave = _clave(nombre)
            if clave in nombres:
                resultado.append(_fila(numero, nombre, "existe", "Ya está en tu catálogo"))
                continue
            if rfc and rfc in rfcs:
                resultado.append(_fila(numero, nombre, "existe", "Ya hay un proveedor con ese RFC"))
                continue
            if clave in vistos_nombre or (rfc and rfc in vistos_rfc):
                previa = vistos_nombre.get(clave) or vistos_rfc.get(rfc)
                resultado.append(_fila(numero, nombre, "repetida", f"Repetido en el archivo (fila {previa})"))
                continue
            correo = (d.get("correo") or "").strip()
            if correo and not CORREO_REGEX.match(correo):
                avisos.append(f"Correo «{correo}» no válido: se omitió")
                correo = ""
            cp = re.sub(r"\D", "", d.get("codigo_postal") or "")
            if d.get("codigo_postal") and len(cp) != 5:
                avisos.append(f"C.P. «{d['codigo_postal']}» no tiene 5 dígitos: se omitió")
                cp = ""
            entrega = Decimal("0")
            if d.get("tiempo_entrega"):
                n = _numero(d["tiempo_entrega"])
                if n is None:
                    avisos.append(f"Días de entrega «{d['tiempo_entrega']}» no es un número: se omitió")
                else:
                    entrega = n
            datos = {
                "nombre": nombre,
                "nombre_comercial": _recortar(d.get("nombre_comercial"), 180, "Nombre comercial", avisos),
                "rfc": rfc,
                "correo": _recortar(correo, 150, "Correo", avisos),
                "telefono": _recortar(d.get("telefono"), 25, "Teléfono", avisos),
                "contacto": _recortar(d.get("contacto"), 150, "Contacto", avisos),
                "direccion": _recortar(d.get("direccion"), 200, "Dirección", avisos),
                "colonia": _recortar(d.get("colonia"), 100, "Colonia", avisos),
                "ciudad": _recortar(d.get("ciudad"), 100, "Ciudad", avisos),
                "codigo_postal": cp,
                "condicion_pago": _recortar(d.get("condicion_pago"), 50, "Condición de pago", avisos),
                "tiempo_entrega": str(entrega),
                "tipo": _recortar(d.get("tipo"), 50, "Tipo", avisos) or "Materia prima",
                "notas": (d.get("notas") or "").strip(),
            }
            vistos_nombre[clave] = numero
            if rfc:
                vistos_rfc[rfc] = numero
            resultado.append(_fila(numero, nombre, "nuevo", "", avisos, datos))
        return resultado

    def crear_proveedores(filas):
        usuario = usuario_actual()
        numero = int(siguiente_codigo().split("-")[1])
        for f in filas:
            d = f["datos"]
            codigo = f"PROV-{numero:04d}"
            numero += 1
            proveedor = Proveedor(
                codigo_proveedor=codigo,
                razon_social=d["nombre"],
                nombre_comercial=d["nombre_comercial"] or None,
                rfc=d["rfc"] or rfc_marcador(codigo),
                tipo_proveedor=d["tipo"],
                estatus="Borrador",
                calle_numero=d["direccion"],
                colonia=d["colonia"],
                ciudad=d["ciudad"] or "Cancún",
                codigo_postal=d["codigo_postal"],
                condicion_pago=d["condicion_pago"],
                tiempo_entrega_normal=Decimal(d["tiempo_entrega"]),
                notas_generales=d["notas"] or None,
                creado_por=usuario,
                actualizado_por=usuario,
            )
            db.session.add(proveedor)
            db.session.flush()
            if d["telefono"] or d["correo"] or d["contacto"]:
                db.session.add(Contacto(
                    id_proveedor=proveedor.id_proveedor,
                    nombre=d["contacto"] or d["nombre"],
                    telefono=d["telefono"] or None,
                    correo=d["correo"],
                    tipo_contacto="Comercial Principal",
                    es_principal=True,
                    activo=True,
                ))
            db.session.add(Bitacora(
                id_proveedor=proveedor.id_proveedor,
                accion="Alta",
                descripcion="Importado desde archivo (Borrador)",
                usuario=usuario,
            ))

    def analizar_clientes(filas):
        por_nombre = {}
        telefonos = set()
        for c in ClienteCRM.query.all():
            por_nombre.setdefault(_clave(c.nombre_completo), []).append(c.telefono or "")
            if c.telefono:
                telefonos.add(_telefono_clave(c.telefono))
        vistos_nombre, vistos_tel = {}, {}
        resultado = []
        for numero, d in filas:
            avisos = []
            nombre = _recortar(d.get("nombre"), 180, "Nombre", avisos)
            if not nombre:
                resultado.append(_fila(numero, "", "omitida", "Sin nombre"))
                continue
            apellidos = _recortar(d.get("apellidos"), 180, "Apellidos", avisos)
            completo = f"{nombre} {apellidos}".strip()
            telefono = _recortar(d.get("telefono"), 40, "Teléfono", avisos)
            tel_clave = _telefono_clave(telefono) if telefono else ""
            clave = _clave(completo)
            if tel_clave and tel_clave in telefonos:
                resultado.append(_fila(numero, completo, "existe", "Ya hay un cliente con ese teléfono"))
                continue
            if clave in por_nombre and (not tel_clave or "" in por_nombre[clave]):
                resultado.append(_fila(numero, completo, "existe", "Ya está en tu directorio"))
                continue
            previa = vistos_tel.get(tel_clave) if tel_clave else None
            previa = previa or (vistos_nombre.get(clave) if not tel_clave else None)
            if previa:
                resultado.append(_fila(numero, completo, "repetida", f"Repetido en el archivo (fila {previa})"))
                continue
            rfc = re.sub(r"\s+", "", d.get("rfc", "")).upper()
            if rfc and not RFC_REGEX.match(rfc):
                avisos.append(f"RFC «{rfc}» no válido: se importa sin RFC")
                rfc = ""
            correo = (d.get("correo") or "").strip()
            if correo and not CORREO_REGEX.match(correo):
                avisos.append(f"Correo «{correo}» no válido: se omitió")
                correo = ""
            direccion = ", ".join(x for x in ((d.get("direccion") or "").strip(), (d.get("colonia") or "").strip()) if x)
            datos = {
                "nombre": nombre,
                "apellidos": apellidos,
                "telefono": telefono,
                "correo": _recortar(correo, 180, "Correo", avisos),
                "rfc": rfc,
                "direccion": _recortar(direccion, 280, "Dirección", avisos),
                "referencias": _recortar(d.get("referencias"), 500, "Referencias", avisos),
                "zona": _recortar(d.get("zona"), 100, "Zona", avisos) or "Cancun",
            }
            vistos_nombre[clave] = numero
            if tel_clave:
                vistos_tel[tel_clave] = numero
            resultado.append(_fila(numero, completo, "nuevo", "", avisos, datos))
        return resultado

    def crear_clientes(filas):
        for f in filas:
            d = f["datos"]
            cliente = ClienteCRM(
                nombre=d["nombre"],
                apellidos=d["apellidos"] or None,
                rfc=d["rfc"] or None,
                telefono=d["telefono"],
                correo=d["correo"] or None,
                canal_origen="Importado",
            )
            db.session.add(cliente)
            db.session.flush()
            if d["direccion"]:
                db.session.add(ClienteDomicilio(
                    cliente_id=cliente.id,
                    alias="Principal",
                    direccion=d["direccion"],
                    referencias=d["referencias"] or None,
                    zona=d["zona"],
                    cobertura=True,
                ))

    TIPOS = {
        "proveedores": (analizar_proveedores, crear_proveedores),
        "clientes": (analizar_clientes, crear_clientes),
    }

    @app.route("/admin/importar")
    def vista_importar():
        if not session.get("admin_logueado"):
            return redirect(url_for("login_admin"))
        tipo = request.args.get("tipo", "proveedores")
        return render_template("importar.html", tipo=tipo if tipo in TIPOS else "proveedores")

    @app.route("/admin/importar/plantilla/<tipo>.xlsx")
    def plantilla_importar(tipo):
        if not session.get("admin_logueado"):
            return redirect(url_for("login_admin"))
        if tipo not in TIPOS:
            return jsonify({"error": "Tipo no válido"}), 404
        return send_file(
            _plantilla_xlsx(tipo),
            as_attachment=True,
            download_name=f"plantilla_{tipo}_nyds.xlsx",
            mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        )

    @app.route("/api/importar/<tipo>", methods=["POST"])
    @admin_required_json
    def api_importar(tipo):
        if tipo not in TIPOS:
            return jsonify({"error": "Elige si importas proveedores o clientes."}), 404
        archivo = request.files.get("archivo")
        if not archivo or not archivo.filename:
            return jsonify({"error": "Elige un archivo de Excel (.xlsx) o CSV."}), 400
        analizar, crear = TIPOS[tipo]
        try:
            leidas, ignoradas = leer_archivo(archivo.filename, archivo.read(MAX_BYTES + 1), tipo)
            filas = analizar(leidas)
        except ErrorImportar as e:
            return jsonify({"error": str(e)}), 400
        resumen = _resumen(filas, ignoradas)
        if not resumen["total"]:
            return jsonify({"error": "El archivo no trae filas con datos debajo del encabezado."}), 400
        if request.form.get("confirmar") == "1":
            nuevos = [f for f in filas if f["estado"] == "nuevo"]
            try:
                crear(nuevos)
                db.session.commit()
            except Exception:
                db.session.rollback()
                return jsonify({"error": "No se pudo guardar. No se importó nada; revisa el archivo e inténtalo de nuevo."}), 400
            return jsonify({"importados": len(nuevos), "resumen": resumen}), 201
        vista = [{k: f[k] for k in ("fila", "nombre", "estado", "motivo", "avisos")} for f in filas[:MAX_FILAS_VISTA]]
        return jsonify({"resumen": resumen, "filas": vista, "recortado": max(0, len(filas) - MAX_FILAS_VISTA)}), 200
