"""Importar proveedores y clientes desde Excel/CSV: solo el nombre es obligatorio,
sin duplicados, datos invalidos se omiten con aviso. SQLite en memoria."""
import io
import unittest
from unittest import mock

from openpyxl import Workbook, load_workbook

from test_proveedores_mp import app, db
from proveedores.vendssback import Bitacora, ClienteCRM, ClienteDomicilio, Contacto, Proveedor, rfc_visible


def xlsx(filas):
    libro = Workbook()
    hoja = libro.active
    for fila in filas:
        hoja.append(fila)
    salida = io.BytesIO()
    libro.save(salida)
    salida.seek(0)
    return salida


def csv_bytes(texto, codificacion="utf-8-sig"):
    return io.BytesIO(texto.encode(codificacion))


class ImportarDatosTest(unittest.TestCase):
    def setUp(self):
        self.ctx = app.app_context()
        self.ctx.push()
        db.drop_all()
        db.create_all()
        self.client = app.test_client()
        with self.client.session_transaction() as sesion:
            sesion["admin_logueado"] = True
            sesion["usuario"] = "compras.nyds"

    def tearDown(self):
        db.session.remove()
        self.ctx.pop()

    def enviar(self, tipo, archivo, nombre="datos.xlsx", confirmar=False):
        datos = {"archivo": (io.BytesIO(archivo.getvalue()), nombre)}  # copia: el archivo se puede reenviar
        if confirmar:
            datos["confirmar"] = "1"
        return self.client.post(f"/api/importar/{tipo}", data=datos, content_type="multipart/form-data")

    # ---- proveedores ----------------------------------------------------
    def test_proveedores_solo_con_nombre_y_vista_previa_no_guarda(self):
        archivo = xlsx([["Nombre"], ["Química del Sureste"], ["Envases Cancún"]])
        r = self.enviar("proveedores", archivo)
        self.assertEqual(r.status_code, 200)
        self.assertEqual(r.get_json()["resumen"]["nuevos"], 2)
        self.assertEqual(Proveedor.query.count(), 0)  # revisar no guarda nada

    def test_proveedores_importa_borrador_con_contacto_y_bitacora(self):
        archivo = xlsx([
            ["Razón social", "Teléfono", "Correo", "RFC", "Contacto", "Dirección", "Colonia", "C.P.", "Días de entrega", "Condición de pago"],
            ["Química del Sureste", 9981234567, "ventas@quimica.mx", "aaa010101aaa", "Ana Pérez", "Av. Tulum 1", "Centro", 77500, "3 días", "Contado"],
            ["Envases Cancún", None, None, None, None, None, None, None, None, None],
        ])
        r = self.enviar("proveedores", archivo, confirmar=True)
        self.assertEqual(r.status_code, 201)
        self.assertEqual(r.get_json()["importados"], 2)
        q = Proveedor.query.filter_by(razon_social="Química del Sureste").one()
        self.assertEqual((q.estatus, q.rfc, q.codigo_postal, float(q.tiempo_entrega_normal)), ("Borrador", "AAA010101AAA", "77500", 3.0))
        contacto = q.contacto_principal()
        self.assertEqual((contacto.nombre, contacto.telefono, contacto.correo), ("Ana Pérez", "9981234567", "ventas@quimica.mx"))
        e = Proveedor.query.filter_by(razon_social="Envases Cancún").one()
        self.assertEqual(rfc_visible(e.rfc), "")  # sin RFC: marcador interno oculto
        self.assertEqual(e.contactos, [])
        self.assertNotEqual(q.codigo_proveedor, e.codigo_proveedor)
        self.assertEqual(Bitacora.query.filter_by(accion="Alta").count(), 2)

    def test_proveedores_no_duplica_existentes_ni_repetidos(self):
        self.enviar("proveedores", xlsx([["Nombre", "RFC"], ["Pochteca", "AAA010101AAA"]]), confirmar=True)
        archivo = xlsx([
            ["Nombre", "RFC"],
            ["  POCHTECA ", ""],            # mismo nombre sin importar mayusculas/espacios
            ["Otro Nombre", "AAA010101AAA"],  # mismo RFC
            ["Nuevo Uno", ""],
            ["nuevo uno", ""],              # repetido en el archivo
        ])
        r = self.enviar("proveedores", archivo).get_json()
        estados = [f["estado"] for f in r["filas"]]
        self.assertEqual(estados, ["existe", "existe", "nuevo", "repetida"])
        self.enviar("proveedores", xlsx([["Nombre"], ["Pochteca"], ["Nuevo Uno"]]), confirmar=True)
        self.assertEqual(Proveedor.query.count(), 2)

    def test_proveedores_datos_invalidos_se_omiten_con_aviso(self):
        archivo = xlsx([
            ["Nombre", "RFC", "Correo", "C.P.", "Días de entrega"],
            ["Proveedor Raro", "123", "no-es-correo", "775", "pronto"],
        ])
        fila = self.enviar("proveedores", archivo).get_json()["filas"][0]
        self.assertEqual(fila["estado"], "nuevo")
        self.assertEqual(len(fila["avisos"]), 4)
        self.enviar("proveedores", archivo, confirmar=True)
        p = Proveedor.query.one()
        self.assertEqual((rfc_visible(p.rfc), p.codigo_postal, float(p.tiempo_entrega_normal)), ("", "", 0.0))
        self.assertEqual(p.contactos, [])

    # ---- clientes -------------------------------------------------------
    def test_clientes_importa_con_domicilio_opcional(self):
        archivo = xlsx([
            ["Nombre", "Apellidos", "Teléfono", "Dirección", "Colonia", "Zona"],
            ["María", "López", "998 765 4321", "Calle 12", "SM 200", "Cancun"],
            ["Pedro", None, None, None, None, None],
        ])
        r = self.enviar("clientes", archivo, confirmar=True)
        self.assertEqual(r.get_json()["importados"], 2)
        maria = ClienteCRM.query.filter_by(nombre="María").one()
        self.assertEqual((maria.apellidos, maria.telefono, maria.canal_origen), ("López", "998 765 4321", "Importado"))
        self.assertEqual(maria.domicilios[0].direccion, "Calle 12, SM 200")
        pedro = ClienteCRM.query.filter_by(nombre="Pedro").one()
        self.assertEqual((pedro.telefono, pedro.domicilios), ("", []))

    def test_clientes_no_duplica_por_telefono_ni_por_nombre(self):
        self.enviar("clientes", xlsx([["Nombre", "Teléfono"], ["Ana", "998 111 2222"], ["Luis", None]]), confirmar=True)
        archivo = xlsx([
            ["Nombre", "Teléfono"],
            ["Otra Persona", "(998) 111-2222"],  # mismo telefono con otro formato
            ["Luis", None],                        # mismo nombre sin telefono
            ["Ana", "998 333 4444"],               # mismo nombre, telefono distinto: es otro cliente
            ["Carla", None],
            ["Carla", None],                       # repetida en el archivo
        ])
        estados = [f["estado"] for f in self.enviar("clientes", archivo).get_json()["filas"]]
        self.assertEqual(estados, ["existe", "existe", "nuevo", "nuevo", "repetida"])

    # ---- archivo --------------------------------------------------------
    def test_csv_con_punto_y_coma_y_acentos_de_excel(self):
        texto = "Nombre;Teléfono;Correo\nJosé Núñez;9981112233;jose@correo.com\n"
        r = self.enviar("clientes", csv_bytes(texto, "cp1252"), nombre="clientes.csv", confirmar=True)
        self.assertEqual(r.status_code, 201)
        self.assertEqual(ClienteCRM.query.one().nombre, "José Núñez")

    def test_filas_sin_nombre_se_omiten_y_vacias_se_ignoran(self):
        archivo = xlsx([["Nombre", "Teléfono"], [None, "998"], [None, None], ["Con Nombre", None]])
        r = self.enviar("clientes", archivo).get_json()
        self.assertEqual([f["estado"] for f in r["filas"]], ["omitida", "nuevo"])
        self.assertEqual(r["resumen"]["omitidas"], 1)

    def test_encabezado_no_tiene_que_ser_la_primera_fila_y_avisa_columnas_ignoradas(self):
        archivo = xlsx([["Lista de proveedores 2026"], [], ["Proveedor", "Saldo"], ["Pochteca", 100]])
        r = self.enviar("proveedores", archivo).get_json()
        self.assertEqual(r["resumen"]["nuevos"], 1)
        self.assertEqual(r["resumen"]["columnas_ignoradas"], ["Saldo"])

    def test_errores_de_archivo_son_claros(self):
        sin_nombre = self.enviar("proveedores", xlsx([["Telefono"], ["998"]]))
        self.assertEqual(sin_nombre.status_code, 400)
        self.assertIn("Nombre", sin_nombre.get_json()["error"])
        self.assertEqual(self.enviar("proveedores", io.BytesIO(b"no es excel"), nombre="x.xlsx").status_code, 400)
        self.assertIn(".xlsx", self.enviar("proveedores", io.BytesIO(b"x"), nombre="viejo.xls").get_json()["error"])
        self.assertEqual(self.enviar("proveedores", io.BytesIO(b"x"), nombre="foto.png").status_code, 400)
        self.assertEqual(self.enviar("otro", xlsx([["Nombre"]])).status_code, 404)
        self.assertEqual(self.client.post("/api/importar/proveedores").status_code, 400)

    def test_si_falla_al_guardar_no_se_importa_nada(self):
        archivo = xlsx([["Nombre"], ["Uno"], ["Dos"]])
        with mock.patch.object(db.session, "commit", side_effect=RuntimeError("falla")):
            r = self.enviar("proveedores", archivo, confirmar=True)
        self.assertEqual(r.status_code, 400)
        self.assertEqual(Proveedor.query.count(), 0)

    # ---- acceso y plantilla ---------------------------------------------
    def test_requiere_admin(self):
        anonimo = app.test_client()
        self.assertEqual(anonimo.post("/api/importar/proveedores").status_code, 403)
        self.assertEqual(anonimo.get("/admin/importar").status_code, 302)
        self.assertEqual(anonimo.get("/admin/importar/plantilla/clientes.xlsx").status_code, 302)

    def test_pantalla_y_plantillas(self):
        self.assertEqual(self.client.get("/admin/importar?tipo=clientes").status_code, 200)
        self.assertIn("Importar desde Excel", self.client.get("/admin/catalogo").get_data(as_text=True))
        for tipo in ("proveedores", "clientes"):
            r = self.client.get(f"/admin/importar/plantilla/{tipo}.xlsx")
            self.assertEqual(r.status_code, 200)
            libro = load_workbook(io.BytesIO(r.data))
            self.assertEqual(libro.sheetnames, ["Datos", "Ejemplo", "Instrucciones"])
            self.assertEqual(libro["Datos"]["A1"].value, "Nombre")
            # la plantilla vacia se acepta como archivo valido de encabezados, pero sin filas no hay nada que importar
            r2 = self.enviar(tipo, io.BytesIO(r.data))
            self.assertEqual(r2.status_code, 400)
        self.assertEqual(self.client.get("/admin/importar/plantilla/otro.xlsx").status_code, 404)

    def test_plantilla_llena_se_importa(self):
        libro = load_workbook(io.BytesIO(self.client.get("/admin/importar/plantilla/proveedores.xlsx").data))
        libro["Datos"].append(["Proveedor de Plantilla", "998 000 1111"])
        salida = io.BytesIO()
        libro.save(salida)
        salida.seek(0)
        self.assertEqual(self.enviar("proveedores", salida, confirmar=True).get_json()["importados"], 1)


if __name__ == "__main__":
    unittest.main()
